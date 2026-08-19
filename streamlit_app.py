"""
Веб-приложение AI-агента ценообразования сети супермаркетов «Зеленое Яблоко».

AI-агент рассчитывает цены, наценку и маржу по отношению к конкурентам
в различных сценариях, формирует итоговый Excel-файл и выводит ключевой
BI-дашборд с выводами прямо в браузере.

Запуск локально:
    python -m streamlit run streamlit_app.py

Развёртывание (Streamlit Community Cloud):
    Укажите этот файл (streamlit_app.py) как главный модуль,
    зависимости берутся из requirements.txt.

Приложение оборачивает существующий публичный API
(za_price_calculator.service.ZAPriceCalculator) без изменения логики модуля.
Метрики дашборда считаются в pandas по той же методике, что и в Excel-книге
(наценка, маржа, отклонение от рынка, сценарии С2 и С4 — релиз 3.7).
"""
from __future__ import annotations

import io
import logging
import tempfile
from pathlib import Path

import altair as alt
import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from za_price_calculator.config import COMPETITOR_COLUMNS, THRESHOLDS
from za_price_calculator.exceptions import ZAPriceCalculatorError
from za_price_calculator.io_handlers.loader import load_sales_file, load_source_file
from za_price_calculator.service import ZAPriceCalculator

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
)

st.set_page_config(
    page_title="Зеленое Яблоко · AI-агент ценообразования",
    page_icon="🍏",
    layout="wide",
    initial_sidebar_state="expanded",
)

# --- Названия конкурентов для человекочитаемых подписей ---
COMPETITOR_LABELS = {
    "Цена_7_Континент": "7 Континент",
    "Цена_Европейский": "Европейский",
    "Цена_Тропики": "Тропики",
    "Цена_Мята": "Мята",
    "Цена_Отличный": "Отличный",
    "Цена_Оптовик_Молоток": "Оптовик (Молоток)",
    "Цена_Оптовик_Редукторный": "Оптовик (Редукторный)",
}

APP_VERSION = "2026-08-19.7"
RELEASE_LABEL = "Релиз 3.7"

GREEN = "#1a7f37"
GREEN_LIGHT = "#2ea043"
RED = "#c0392b"
AMBER = "#d4a017"
BLUE = "#1f6feb"
ORANGE = "#ed7d31"
SKY = "#5b9bd5"
DEV_HIGHLIGHT = 0.15  # 15% — порог подсветки отклонений

# --- Шаблон файла продаж ---
SALES_TEMPLATE_HEADERS = ["Штрихкод", "Кол-во продаж", "Выручка", "Валовая прибыль"]
# Примеры используют реальные штрихкоды из прайса — при загрузке они сопоставятся с товарами.
SALES_TEMPLATE_ROWS = [
    ("4665272570017", 120, 3480, 900),
    ("4602984001637", 85, 8330, 2100),
    ("8690511171614", 40, 25960, 5200),
]


@st.cache_data(show_spinner=False)
def _sales_template_bytes() -> bytes:
    """Формирует Excel-шаблон файла продаж с заголовками и примерами строк."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Продажи"

    header_fill = PatternFill("solid", fgColor="1A7F37")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    for ci, title in enumerate(SALES_TEMPLATE_HEADERS, start=1):
        cell = ws.cell(1, ci, title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for ri, row in enumerate(SALES_TEMPLATE_ROWS, start=2):
        for ci, value in enumerate(row, start=1):
            cell = ws.cell(ri, ci, value)
            if ci == 1:  # штрихкод храним как текст, чтобы не терять ведущие нули
                cell.number_format = "@"

    for ci, width in enumerate((22, 16, 14, 18), start=1):
        ws.column_dimensions[get_column_letter(ci)].width = width
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

# =============================== Оформление =================================
st.markdown(
    """
    <style>
      #MainMenu, footer {visibility: hidden;}
      .block-container {padding-top: 1.4rem; padding-bottom: 2rem; max-width: 1280px;}

      .za-hero {
        background: linear-gradient(125deg, #0b3d28 0%, #147a3a 42%, #2bb34f 100%);
        border-radius: 20px; padding: 28px 34px; color: #fff;
        box-shadow: 0 14px 36px rgba(15,81,50,.35); margin-bottom: 10px;
      }
      .za-badge {
        display: inline-block;
        background: linear-gradient(90deg, #ffe08a 0%, #fff3c4 100%);
        color: #0f5132; font-size: 22px; font-weight: 800; letter-spacing: .2px;
        padding: 10px 18px; border-radius: 14px; margin-bottom: 16px;
        box-shadow: 0 6px 18px rgba(0,0,0,.18); border: 2px solid rgba(255,255,255,.55);
      }
      .za-hero h1 {
        margin: 0 0 12px 0; font-size: 38px; font-weight: 900; letter-spacing: .3px;
        text-shadow: 0 2px 10px rgba(0,0,0,.18);
      }
      .za-hero p {
        margin: 0; font-size: 18px; font-weight: 600; line-height: 1.55;
        color: #ffffff; max-width: 980px; text-shadow: 0 1px 4px rgba(0,0,0,.15);
      }
      .za-hero p b { color: #ffe08a; font-weight: 800; }

      .za-loading {
        text-align: center; padding: 48px 20px; margin: 18px 0 10px;
        background: linear-gradient(180deg, #e8f7ee 0%, #ffffff 100%);
        border: 2px solid #b7e0c4; border-radius: 22px;
        box-shadow: 0 10px 28px rgba(26,127,55,.14);
      }
      .za-loading .bot { font-size: 92px; line-height: 1; margin-bottom: 8px; }
      .za-loading .title {
        font-size: 34px; font-weight: 900; color: #0f5132; letter-spacing: .3px;
      }
      .za-loading .sub { margin-top: 8px; font-size: 16px; font-weight: 600; color: #3d6b52; }

      div[data-testid="stMetric"] {
        background: linear-gradient(180deg, #ffffff 0%, #f3fbf6 100%);
        border: 1px solid #cfe8d8; border-radius: 18px;
        padding: 18px 18px 14px; box-shadow: 0 8px 22px rgba(16,24,40,.10);
        min-height: 118px;
      }
      div[data-testid="stMetricLabel"] p {
        font-size: 15px !important; color:#1f3d2d !important; font-weight:800 !important;
      }
      div[data-testid="stMetricValue"] {
        font-size: 34px !important; color:#0f5132 !important; font-weight: 900 !important;
      }
      div[data-testid="stMetricDelta"] { font-size: 14px !important; font-weight: 700 !important; }

      .za-section-title {
        font-size: 24px; font-weight: 900; color: #0f5132; margin: 6px 0 12px;
      }
      .za-chart-title {
        font-size: 20px; font-weight: 800; color: #123d28; margin: 0 0 10px;
      }
      .za-table-title {
        font-size: 17px; font-weight: 800; color: #123d28; margin: 10px 0 8px;
      }

      .za-insight {
        border-radius: 16px; padding: 18px 20px; margin-bottom: 12px;
        font-size: 17px; font-weight: 650; line-height: 1.55;
        border: 1px solid transparent; box-shadow: 0 8px 20px rgba(16,24,40,.10);
      }
      .za-insight b { font-weight: 900; }
      .za-ok   { background: #e5f6ea; border-color:#8fd1a4; color:#0f5132; border-left: 8px solid #1a7f37; }
      .za-warn { background: #fff4d6; border-color:#f0d27a; color:#6a4b00; border-left: 8px solid #d4a017; }
      .za-bad  { background: #fde8e6; border-color:#f0a8a0; color:#7a1f18; border-left: 8px solid #c0392b; }
      .za-info { background: #e8f1ff; border-color:#a8c4f0; color:#12325f; border-left: 8px solid #1f6feb; }

      .za-html-table { width:100%; border-collapse: separate; border-spacing: 0; margin: 0 0 14px;
        border-radius: 14px; overflow: hidden; box-shadow: 0 8px 20px rgba(16,24,40,.08);
        border: 1px solid #d7e5dc; }
      .za-html-table th {
        background: #0f5132; color: #fff; font-size: 14px; font-weight: 800;
        padding: 11px 12px; text-align: left;
      }
      .za-html-table td {
        font-size: 14.5px; font-weight: 700; color: #1a2e24; padding: 10px 12px;
        border-top: 1px solid #e4eee8; background: #fff;
      }
      .za-html-table tr:nth-child(even) td { background: #f7fbf8; }
      .za-cell-orange { background: #ffe0c2 !important; color: #8a3b00 !important; }
      .za-cell-sky { background: #d7ebff !important; color: #0b3d66 !important; }

      .za-scen-table { width:100%; border-collapse: separate; border-spacing: 0; margin: 8px 0 16px;
        border-radius: 16px; overflow: hidden; box-shadow: 0 10px 26px rgba(16,24,40,.12);
        border: 2px solid #1a7f37; }
      .za-scen-table th {
        background: linear-gradient(90deg, #0f5132, #1a7f37); color:#fff;
        font-size: 15px; font-weight: 900; padding: 13px 14px; text-align: left;
      }
      .za-scen-table td {
        font-size: 15.5px; font-weight: 750; color:#10281c; padding: 12px 14px;
        border-top: 1px solid #cfe8d8; background: #f2faf5;
      }
      .za-scen-table tr:nth-child(even) td { background: #e7f6ec; }
      .za-scen-table tr:hover td { background: #d7f0e0; }

      .stTabs [data-baseweb="tab-list"] { gap: 8px; }
      .stTabs [data-baseweb="tab"] { font-weight: 800; font-size: 15px; }
    </style>
    """,
    unsafe_allow_html=True,
)

st.markdown(
    f"""
    <div class="za-hero">
      <div class="za-badge">🤖 AI-агент ценообразования · {RELEASE_LABEL}</div>
      <h1>🍏 Зеленое Яблоко</h1>
      <p>Интеллектуальный агент по расчёту <b>цен, наценки и маржи</b> сети супермаркетов
      «Зеленое Яблоко» по отношению к <b>конкурентам</b> в различных
      <b>сценариях</b>. Загрузите прайс-лист — агент сформирует готовый Excel-файл
      и покажет ключевой BI-дашборд с выводами и рекомендациями.</p>
    </div>
    """,
    unsafe_allow_html=True,
)

st.caption(
    f"🔖 {RELEASE_LABEL} · версия интерфейса: {APP_VERSION} · "
    "обновлённый UI дашборда, таблицы отклонений, сценарии С2/С4. "
    "Если подпись «Релиз 3.7» не видна — на Streamlit Cloud: Manage app → Reboot, "
    "затем Ctrl/Cmd+Shift+R."
)

with st.expander("ℹ️ Что нового в релизе 3.7", expanded=False):
    st.markdown(
        """
**Визуал Streamlit**
- Крупный яркий заголовок AI-агента и выразительное описание в шапке
- Экран загрузки: крупный 🤖 и текст «AI агент работает»
- Объёмные KPI-карточки, крупные подписи к графикам
- ТОП отклонений: добавлена цена конкурента; подсветка ±15%
- Выводы AI — светофорные блоки; сценарии — насыщенная таблица
- При наличии продаж: текущая ВП и прирост при выравнивании до медианы
  (только позиции с ценой ниже медианы рынка)
        """
    )


# =============================== Логика расчётов =============================
@st.cache_data(show_spinner=False)
def _load_source_df(file_bytes: bytes, file_name: str) -> pd.DataFrame:
    """Загружает и нормализует прайс-лист во временном файле (кэшируется по байтам)."""
    suffix = Path(file_name).suffix or ".xlsx"
    with tempfile.NamedTemporaryFile(suffix=suffix, delete=True) as tmp:
        tmp.write(file_bytes)
        tmp.flush()
        return load_source_file(tmp.name)


@st.cache_data(show_spinner=False)
def _load_sales_df(file_bytes: bytes, file_name: str) -> pd.DataFrame:
    """Загружает и нормализует файл продаж во временном файле."""
    suffix = Path(file_name).suffix or ".xlsx"
    with tempfile.NamedTemporaryFile(suffix=suffix, delete=True) as tmp:
        tmp.write(file_bytes)
        tmp.flush()
        return load_sales_file(tmp.name)


@st.cache_data(show_spinner=False)
def _build_excel(
    source_bytes: bytes, source_name: str,
    sales_bytes: bytes | None, sales_name: str | None,
) -> bytes:
    """Формирует итоговый Excel-файл через публичный API и возвращает его байты."""
    with tempfile.TemporaryDirectory() as tmp:
        tmp_dir = Path(tmp)
        src = tmp_dir / (Path(source_name).stem + (Path(source_name).suffix or ".xlsx"))
        src.write_bytes(source_bytes)
        sales = None
        if sales_bytes is not None and sales_name:
            sales = tmp_dir / (Path(sales_name).stem + (Path(sales_name).suffix or ".xlsx"))
            sales.write_bytes(sales_bytes)
        out = tmp_dir / "Зеленое_Яблоко_калькулятор.xlsx"
        ZAPriceCalculator().run(
            source_path=str(src),
            output_path=str(out),
            sales_path=str(sales) if sales else None,
        )
        return out.read_bytes()


def _merge_sales(metrics: pd.DataFrame, sales_df: pd.DataFrame | None) -> pd.DataFrame:
    """Добавляет к метрикам кол-во продаж / выручку / ВП из файла продаж по штрихкоду."""
    d = metrics.copy()
    if sales_df is None or sales_df.empty:
        d["Кол-во_продаж"] = pd.NA
        d["Выручка_продаж"] = pd.NA
        d["ВП_продаж"] = pd.NA
        d["Есть_продажи"] = False
        return d

    sales = sales_df.copy()
    sales["Штрихкод"] = sales["Штрихкод"].astype(str)
    # При дубликатах штрихкода суммируем продажи
    agg = (
        sales.groupby("Штрихкод", as_index=False)
        .agg({"Кол-во_продаж": "sum", "Выручка": "sum", "Валовая_прибыль": "sum"})
        .rename(columns={"Выручка": "Выручка_продаж", "Валовая_прибыль": "ВП_продаж"})
    )
    d["Штрихкод"] = d["Штрихкод"].astype(str)
    d = d.merge(agg, on="Штрихкод", how="left")
    d["Есть_продажи"] = d["Кол-во_продаж"].notna() & (d["Кол-во_продаж"] > 0)
    return d

def _compute_metrics(df: pd.DataFrame) -> pd.DataFrame:
    """Считает наценку, маржу, отклонение от рынка и сигналы по методике модуля."""
    d = df.copy()
    # Приводим цены конкурентов к числам: пустые/нечисловые столбцы иначе остаются
    # object, из-за чего Ср_рынок/Откл_рынок становятся object и ломают nlargest.
    comp = d[COMPETITOR_COLUMNS].apply(pd.to_numeric, errors="coerce")
    d[COMPETITOR_COLUMNS] = comp
    d["Ср_рынок"] = comp.mean(axis=1)
    d["Медиана_рынок"] = comp.median(axis=1)
    d["Мин_конкурент"] = comp.min(axis=1)

    purch = pd.to_numeric(d["Закупочная_цена"], errors="coerce")
    za = pd.to_numeric(d["Розничная_цена_ЗЯ"], errors="coerce")
    d["Закупочная_цена"] = purch
    d["Розничная_цена_ЗЯ"] = za

    d["Валидна"] = purch.notna() & (purch > 0) & za.notna() & (za > 0)
    d["ВП_ед"] = za - purch
    d["Наценка"] = (za - purch) / purch
    d["Маржа"] = (za - purch) / za
    d["Откл_рынок"] = (za - d["Ср_рынок"]) / d["Ср_рынок"]
    d["Откл_медиана"] = (za - d["Медиана_рынок"]) / d["Медиана_рынок"]

    def signal(x):
        if pd.isna(x):
            return "Нет данных рынка"
        if x > THRESHOLDS.price_above_market:
            return "Выше рынка"
        if x < THRESHOLDS.price_below_market:
            return "Ниже рынка"
        return "На уровне"

    d["Сигнал"] = d["Откл_рынок"].apply(signal)

    def margin_class(m):
        if pd.isna(m):
            return "—"
        if m > THRESHOLDS.margin_high:
            return "Хорошая (>20%)"
        if m < THRESHOLDS.margin_low:
            return "Низкая (<10%)"
        return "Средняя (10–20%)"

    d["Класс_маржи"] = d["Маржа"].apply(margin_class)

    # Риск при выравнивании до минимального конкурента (сценарий С3)
    m_c3 = (d["Мин_конкурент"] - purch) / d["Мин_конкурент"]
    d["Маржа_С3"] = m_c3

    def risk(m):
        if pd.isna(m):
            return "—"
        if m < THRESHOLDS.risk_critical:
            return "КРИТИЧНО"
        if m < THRESHOLDS.risk_moderate:
            return "УМЕРЕННЫЙ"
        return "ОК"

    d["Риск_С3"] = m_c3.apply(risk)
    return d


def _scenarios(d: pd.DataFrame) -> pd.DataFrame:
    """Сводка по видимым сценариям (С2, С4) + текущий уровень."""
    v = d[d["Валидна"]].copy()
    purch = v["Закупочная_цена"]

    def block(name, price, note):
        margin = (price - purch) / price
        vp = price - purch
        return {
            "Сценарий": name,
            "Цена": float(price.mean()) if len(price) else float("nan"),
            "Маржа": float(margin.mean()) if len(margin) else float("nan"),
            "ВП_ед": float(vp.mean()) if len(vp) else float("nan"),
            "Вывод": note,
        }

    # С4 в Excel заполняется вручную через «Новая розничная цена ЗЯ»;
    # в веб-превью показываем строку-пояснение без расчётной цены.
    rows = [
        block("Текущая", v["Розничная_цена_ЗЯ"], "Базовый уровень"),
        block("С2 · Медиана рынка", v["Медиана_рынок"], "Устойчив к выбросам"),
        {
            "Сценарий": "С4 · Произв.цена",
            "Цена": float("nan"),
            "Маржа": float("nan"),
            "ВП_ед": float("nan"),
            "Вывод": "Ввод в Excel: «Новая Розничная цена ЗЯ» → С4",
        },
    ]
    return pd.DataFrame(rows)


def _insight(kind: str, text: str) -> None:
    st.markdown(f'<div class="za-insight za-{kind}">{text}</div>', unsafe_allow_html=True)


def _fmt_pct(x) -> str:
    return "—" if pd.isna(x) else f"{x * 100:.1f}%"


def _fmt_num(x) -> str:
    return "—" if pd.isna(x) else f"{x:,.2f}".replace(",", " ")


def _html_escape(s) -> str:
    return (
        str(s)
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
    )


def _render_html_table(headers: list[str], rows: list[list[tuple[str, str | None]]]) -> None:
    """Рендерит HTML-таблицу. Ячейка = (текст, css-класс|None). Без pandas Styler/jinja2."""
    th = "".join(f"<th>{_html_escape(h)}</th>" for h in headers)
    body = []
    for row in rows:
        tds = []
        for text, css in row:
            cls = f' class="{css}"' if css else ""
            tds.append(f"<td{cls}>{_html_escape(text)}</td>")
        body.append("<tr>" + "".join(tds) + "</tr>")
    html = (
        '<table class="za-html-table"><thead><tr>'
        + th
        + "</tr></thead><tbody>"
        + "".join(body)
        + "</tbody></table>"
    )
    st.markdown(html, unsafe_allow_html=True)


def _dev_cell_class(dev: float, above: bool) -> str | None:
    """Подсветка: выше >15% — оранжевый; ниже >15% — голубой."""
    if pd.isna(dev):
        return None
    if above and dev > DEV_HIGHLIGHT:
        return "za-cell-orange"
    if (not above) and dev < -DEV_HIGHLIGHT:
        return "za-cell-sky"
    return None


def _top_dev_rows(
    df: pd.DataFrame, col: str, ref_col: str, ascending: bool, n: int = 10,
) -> list[list[tuple[str, str | None]]]:
    subset = df[df[col].notna()].copy()
    if subset.empty:
        return []
    top = subset.nsmallest(n, col) if ascending else subset.nlargest(n, col)
    above = not ascending
    rows: list[list[tuple[str, str | None]]] = []
    for _, r in top.iterrows():
        dev = r[col]
        pct = "—" if pd.isna(dev) else f"{dev * 100:+.1f}%"
        rows.append([
            (str(r["Наименование"])[:70], None),
            (pct, _dev_cell_class(dev, above=above)),
            (_fmt_num(r["Розничная_цена_ЗЯ"]), None),
            (_fmt_num(r[ref_col]), None),
        ])
    return rows


def _below_median_sales_impact(d: pd.DataFrame) -> dict | None:
    """
    Для позиций с ценой ЗЯ ниже медианы рынка и с продажами:
    текущая ВП и ВП при выравнивании цены до медианы.
    """
    if "Есть_продажи" not in d.columns or not d["Есть_продажи"].any():
        return None
    m = d[
        d["Есть_продажи"]
        & d["Валидна"]
        & d["Медиана_рынок"].notna()
        & (d["Розничная_цена_ЗЯ"] < d["Медиана_рынок"])
    ].copy()
    if m.empty:
        return None
    qty = pd.to_numeric(m["Кол-во_продаж"], errors="coerce").fillna(0)
    # Текущая ВП: из файла продаж, иначе qty × ВП/ед
    vp_file = pd.to_numeric(m.get("ВП_продаж"), errors="coerce")
    vp_calc = qty * pd.to_numeric(m["ВП_ед"], errors="coerce")
    current_vp = vp_file.where(vp_file.notna(), vp_calc).fillna(0)
    # Новая ВП при цене = медиана: qty × (медиана − закупка)
    new_vp_unit = m["Медиана_рынок"] - m["Закупочная_цена"]
    new_vp = qty * new_vp_unit
    cur_sum = float(current_vp.sum())
    new_sum = float(new_vp.sum())
    return {
        "n_pos": int(len(m)),
        "qty": float(qty.sum()),
        "current_vp": cur_sum,
        "median_vp": new_sum,
        "delta_vp": new_sum - cur_sum,
    }


# =============================== Панель загрузки =============================
with st.sidebar:
    st.markdown("### 📂 Загрузка данных")
    source_upload = st.file_uploader(
        "Прайс-лист (обязательно)",
        type=["xlsx", "xls", "xlsm"],
        help="Excel с товарами, закупочными и розничными ценами ЗЯ и ценами конкурентов.",
    )
    sales_upload = st.file_uploader(
        "Файл продаж (опционально)",
        type=["xlsx", "xls", "xlsm"],
        help="Колонки: Штрихкод (обязательно), Кол-во продаж, Выручка, Валовая прибыль. "
             "Строки сопоставляются с прайсом по штрихкоду.",
    )
    st.download_button(
        "📥 Скачать шаблон продаж (.xlsx)",
        data=_sales_template_bytes(),
        file_name="шаблон_продаж_ЗЯ.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        width="stretch",
        help="Готовый шаблон с нужными колонками и примерами строк.",
        key="dl_tmpl_sidebar",
    )
    st.caption(
        f"**{RELEASE_LABEL}.** Файл продаж — блок «Файл продаж»; обязателен **Штрихкод**. "
        "UI дашборда обновлён: крупные KPI, подсветка отклонений ±15%, сценарии С2/С4."
    )

if source_upload is None:
    st.info(
        "⬅️ Загрузите прайс-лист на панели слева, чтобы AI-агент рассчитал цены "
        "и показал дашборд с выводами по сценариям."
    )
    st.markdown("#### С чего начать")
    step1, step2 = st.columns(2)
    with step1:
        st.markdown(
            "**1. Прайс-лист (обязательно)** — загрузите на панели слева "
            "в блоке «Прайс-лист». Excel с ценами ЗЯ и конкурентов."
        )
    with step2:
        st.markdown(
            "**2. Продажи (опционально)** — загрузите в блок «Файл продаж». "
            "Обязательная колонка — **Штрихкод**. Ниже можно скачать готовый шаблон:"
        )
        st.download_button(
            "📥 Скачать шаблон продаж (.xlsx)",
            data=_sales_template_bytes(),
            file_name="шаблон_продаж_ЗЯ.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            help="Готовый шаблон с нужными колонками и примерами строк.",
            key="dl_tmpl_intro",
        )
    st.stop()

# --- Загрузка и расчёт ---
loading_box = st.empty()
loading_box.markdown(
    """
    <div class="za-loading">
      <div class="bot">🤖</div>
      <div class="title">AI агент работает</div>
      <div class="sub">Анализирую прайс-лист, считаю наценку, маржу и сценарии…</div>
    </div>
    """,
    unsafe_allow_html=True,
)
try:
    with st.spinner("AI-агент анализирует данные..."):
        source_bytes = source_upload.getvalue()
        sales_bytes = sales_upload.getvalue() if sales_upload else None
        df = _load_source_df(source_bytes, source_upload.name)
        sales_df = (
            _load_sales_df(sales_bytes, sales_upload.name)
            if sales_upload is not None else None
        )
        metrics = _merge_sales(_compute_metrics(df), sales_df)
        scen = _scenarios(metrics)
        excel_bytes = _build_excel(
            source_bytes, source_upload.name, sales_bytes,
            sales_upload.name if sales_upload else None,
        )
except ZAPriceCalculatorError as exc:
    loading_box.empty()
    st.error(f"Ошибка обработки данных: {exc}")
    st.stop()
except Exception as exc:  # noqa: BLE001
    loading_box.empty()
    st.error("Непредвиденная ошибка при расчёте.")
    st.exception(exc)
    st.stop()
else:
    loading_box.empty()

valid = metrics[metrics["Валидна"]]
n_all = len(metrics)
n_valid = len(valid)
n_sales_matched = int(metrics["Есть_продажи"].sum()) if "Есть_продажи" in metrics.columns else 0
n_sales_rows = len(sales_df) if sales_df is not None else 0

# --- Верхняя строка: статус + скачивание ---
c1, c2 = st.columns([3, 1])
with c1:
    sales_note = ""
    if sales_df is not None:
        sales_note = (
            f" Файл продаж: **{n_sales_rows}** строк, "
            f"сопоставлено с прайсом: **{n_sales_matched}**."
        )
    st.success(
        f"Готово! Обработано позиций: **{n_all}**, с полной экономикой: **{n_valid}**."
        f"{sales_note} Итоговый Excel-файл сформирован."
    )
with c2:
    st.download_button(
        "⬇️ Скачать Excel",
        data=excel_bytes,
        file_name="Зеленое_Яблоко_калькулятор.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        width="stretch",
    )

tab_dash, tab_scen, tab_data = st.tabs(
    ["📊 Дашборд", "🎯 Сценарии", "📋 Данные и выгрузка"]
)

# =============================== Вкладка: Дашборд ============================
with tab_dash:
    avg_margin = valid["Маржа"].mean()
    avg_markup = valid["Наценка"].mean()
    avg_za = valid["Розничная_цена_ЗЯ"].mean()
    avg_vp = valid["ВП_ед"].mean()
    n_above = int((metrics["Сигнал"] == "Выше рынка").sum())
    n_below = int((metrics["Сигнал"] == "Ниже рынка").sum())
    n_level = int((metrics["Сигнал"] == "На уровне").sum())
    n_crit = int((metrics["Риск_С3"] == "КРИТИЧНО").sum())
    n_good = int((metrics["Класс_маржи"] == "Хорошая (>20%)").sum())
    n_low = int((metrics["Класс_маржи"] == "Низкая (<10%)").sum())

    st.markdown('<div class="za-section-title">Ключевые показатели</div>', unsafe_allow_html=True)
    k = st.columns(4)
    k[0].metric("Позиций в прайсе", f"{n_all:,}".replace(",", " "))
    k[1].metric("Средняя маржа", _fmt_pct(avg_margin))
    k[2].metric("Средняя наценка", _fmt_pct(avg_markup))
    k[3].metric("Средняя ВП / ед.", _fmt_num(avg_vp))
    k2 = st.columns(4)
    k2[0].metric("Средняя цена ЗЯ", _fmt_num(avg_za))
    k2[1].metric("Выше рынка", f"{n_above}", delta=f"{n_above / n_all * 100:.0f}%" if n_all else None,
                 delta_color="inverse")
    k2[2].metric("Ниже рынка", f"{n_below}", delta=f"{n_below / n_all * 100:.0f}%" if n_all else None)
    k2[3].metric(
        "Критический риск маржи",
        f"{n_crit}",
        help="Сколько позиций получили бы маржу <5%, если снизить цену ЗЯ "
             "до минимального конкурента. Это не текущий убыток, а оценка запаса прочности.",
        delta_color="inverse",
    )
    st.caption(
        "💡 **Критический риск маржи** — число позиций, у которых маржа упала бы **ниже 5%**, "
        "если выровнять цену ЗЯ до **минимального конкурента**. "
        "КРИТИЧНО (&lt;5%) / УМЕРЕННЫЙ (5–15%) / НИЗКИЙ (&gt;15%). "
        "Это сценарий «что будет при сильном снижении цены», а не текущий убыток."
    )

    st.divider()
    g1, g2 = st.columns(2)

    axis_label = alt.Axis(labelFontSize=14, labelFontWeight="bold", titleFontSize=14, titleFontWeight="bold")
    with g1:
        st.markdown(
            '<div class="za-chart-title">Позиционирование по цене относительно рынка</div>',
            unsafe_allow_html=True,
        )
        sig_df = (
            metrics["Сигнал"].value_counts()
            .rename_axis("Сигнал").reset_index(name="Позиции")
        )
        color_scale = alt.Scale(
            domain=["Ниже рынка", "На уровне", "Выше рынка", "Нет данных рынка"],
            range=[GREEN, BLUE, RED, "#8b949e"],
        )
        chart = (
            alt.Chart(sig_df)
            .mark_bar(cornerRadiusEnd=6)
            .encode(
                x=alt.X("Позиции:Q", title="Кол-во позиций", axis=axis_label),
                y=alt.Y("Сигнал:N", sort="-x", title=None, axis=axis_label),
                color=alt.Color("Сигнал:N", scale=color_scale, legend=None),
                tooltip=["Сигнал", "Позиции"],
            )
            .properties(height=240)
            .configure_axis(labelFontSize=14, titleFontSize=14)
        )
        st.altair_chart(chart, width="stretch")

    with g2:
        st.markdown(
            '<div class="za-chart-title">Распределение по уровню маржи</div>',
            unsafe_allow_html=True,
        )
        mrg_df = (
            valid["Класс_маржи"].value_counts()
            .rename_axis("Класс").reset_index(name="Позиции")
        )
        mrg_scale = alt.Scale(
            domain=["Низкая (<10%)", "Средняя (10–20%)", "Хорошая (>20%)"],
            range=[RED, AMBER, GREEN],
        )
        chart2 = (
            alt.Chart(mrg_df)
            .mark_bar(cornerRadiusEnd=6)
            .encode(
                x=alt.X("Позиции:Q", title="Кол-во позиций", axis=axis_label),
                y=alt.Y("Класс:N", sort="-x", title=None, axis=axis_label),
                color=alt.Color("Класс:N", scale=mrg_scale, legend=None),
                tooltip=["Класс", "Позиции"],
            )
            .properties(height=240)
            .configure_axis(labelFontSize=14, titleFontSize=14)
        )
        st.altair_chart(chart2, width="stretch")

    headers_dev = ["Наименование", "Откл.", "Цена ЗЯ", "Цена конкурента"]

    st.divider()
    st.markdown(
        '<div class="za-section-title">ТОП-10 отклонений от рынка</div>',
        unsafe_allow_html=True,
    )
    st.caption(
        "Оранжевая заливка — отклонение **выше** на &gt;15%. "
        "Голубая заливка — отклонение **ниже** на &gt;15%. "
        "«Цена конкурента» — средняя или медиана рынка (база расчёта)."
    )
    t1, t2 = st.columns(2)
    with t1:
        st.markdown('<div class="za-table-title">Выше рынка по средней</div>', unsafe_allow_html=True)
        rows = _top_dev_rows(metrics, "Откл_рынок", "Ср_рынок", ascending=False)
        if rows:
            _render_html_table(headers_dev, rows)
        else:
            st.caption("Нет данных.")
        st.markdown('<div class="za-table-title">Ниже рынка по средней</div>', unsafe_allow_html=True)
        rows = _top_dev_rows(metrics, "Откл_рынок", "Ср_рынок", ascending=True)
        if rows:
            _render_html_table(headers_dev, rows)
        else:
            st.caption("Нет данных.")
    with t2:
        st.markdown('<div class="za-table-title">Выше рынка по медиане</div>', unsafe_allow_html=True)
        rows = _top_dev_rows(metrics, "Откл_медиана", "Медиана_рынок", ascending=False)
        if rows:
            _render_html_table(headers_dev, rows)
        else:
            st.caption("Нет данных.")
        st.markdown('<div class="za-table-title">Ниже рынка по медиане</div>', unsafe_allow_html=True)
        rows = _top_dev_rows(metrics, "Откл_медиана", "Медиана_рынок", ascending=True)
        if rows:
            _render_html_table(headers_dev, rows)
        else:
            st.caption("Нет данных.")

    st.divider()
    st.markdown('<div class="za-section-title">🧠 Выводы AI-агента</div>', unsafe_allow_html=True)

    if n_all:
        share_above = n_above / n_all * 100
        share_below = n_below / n_all * 100
        if n_crit:
            _insight(
                "bad",
                f"<b>{n_crit}</b> позиций — <b>критический риск маржи</b>: при снижении "
                "до мин.конкурента маржа стала бы &lt;5%. Снижать цену опасно; "
                "используйте С2 или ручной ввод в С4.",
            )
        if n_above:
            _insight(
                "bad",
                f"<b>{n_above}</b> позиций ({share_above:.0f}%) стоят <b>выше рынка</b>. "
                "Рекомендуется рассмотреть сценарий <b>С2 (медиана рынка)</b> "
                "или задать новую розничную цену в колонке «Новая Розничная цена ЗЯ» (С4).",
            )
        if n_low:
            _insight(
                "warn",
                f"<b>{n_low}</b> позиций имеют <b>низкую маржу</b> (&lt;10%). "
                "Приоритет: переговоры с поставщиком или корректировка цены.",
            )
        if n_below:
            _insight(
                "ok",
                f"<b>{n_below}</b> позиций ({share_below:.0f}%) стоят <b>ниже рынка</b>. "
                "Есть потенциал повышения цены и маржи без потери конкурентоспособности.",
            )
        _insight(
            "info",
            f"Средняя маржа по прайсу — <b>{_fmt_pct(avg_margin)}</b>, "
            f"средняя наценка — <b>{_fmt_pct(avg_markup)}</b>. "
            f"Позиций с хорошей маржой (&gt;20%): <b>{n_good}</b> — у них максимальная ценовая гибкость.",
        )

# =============================== Вкладка: Сценарии ==========================
with tab_scen:
    st.markdown(
        '<div class="za-section-title">Сравнение сценариев ценообразования</div>',
        unsafe_allow_html=True,
    )
    st.markdown(
        '<p style="font-size:16px;font-weight:700;color:#1a3d2a;margin:0 0 12px;">'
        "Для пользователя видимы <b>С2: Медиана рынка</b> и "
        "<b>С4: Произв.цена</b>. Сценарии С1, С3, С5 скрыты. "
        "С4 заполняется вручную в колонке «Новая Розничная цена ЗЯ»."
        "</p>",
        unsafe_allow_html=True,
    )

    scen_rows_html = []
    for _, r in scen.iterrows():
        scen_rows_html.append(
            "<tr>"
            f"<td>{_html_escape(r['Сценарий'])}</td>"
            f"<td>{_html_escape(_fmt_num(r['Цена']))}</td>"
            f"<td>{_html_escape(_fmt_pct(r['Маржа']))}</td>"
            f"<td>{_html_escape(_fmt_num(r['ВП_ед']))}</td>"
            f"<td>{_html_escape(r['Вывод'])}</td>"
            "</tr>"
        )
    st.markdown(
        '<table class="za-scen-table"><thead><tr>'
        "<th>Сценарий</th><th>Ср цена</th><th>Маржа</th><th>ВП ед</th><th>Вывод</th>"
        "</tr></thead><tbody>"
        + "".join(scen_rows_html)
        + "</tbody></table>",
        unsafe_allow_html=True,
    )

    st.markdown(
        '<div class="za-chart-title">Средняя маржа по сценариям</div>',
        unsafe_allow_html=True,
    )
    scen_chart_df = scen[scen["Маржа"].notna()].copy()
    if len(scen_chart_df):
        scen_chart = (
            alt.Chart(scen_chart_df)
            .mark_bar(cornerRadiusEnd=6)
            .encode(
                x=alt.X(
                    "Сценарий:N",
                    sort=list(scen_chart_df["Сценарий"]),
                    title=None,
                    axis=alt.Axis(labelAngle=0, labelLimit=180, labelFontSize=14, labelFontWeight="bold"),
                ),
                y=alt.Y(
                    "Маржа:Q",
                    title="Средняя маржа",
                    axis=alt.Axis(format="%", titleFontSize=14, titleFontWeight="bold", labelFontSize=13),
                ),
                color=alt.Color(
                    "Сценарий:N",
                    scale=alt.Scale(
                        domain=list(scen_chart_df["Сценарий"]),
                        range=[BLUE, GREEN],
                    ),
                    legend=None,
                ),
                tooltip=[
                    "Сценарий",
                    alt.Tooltip("Маржа:Q", format=".1%"),
                    alt.Tooltip("ВП_ед:Q", format=".2f"),
                ],
            )
            .properties(height=320)
        )
        st.altair_chart(scen_chart, width="stretch")

    base_m = scen.loc[scen["Сценарий"] == "Текущая", "Маржа"].iloc[0]
    c2_m = scen.loc[scen["Сценарий"] == "С2 · Медиана рынка", "Маржа"].iloc[0]
    impact = _below_median_sales_impact(metrics)

    if impact is None:
        _insight(
            "info",
            f"Текущая средняя маржа <b>{_fmt_pct(base_m)}</b>. При выравнивании до медианы рынка (С2) "
            f"она составит <b>{_fmt_pct(c2_m)}</b>. Сценарий С4 задаётся вручную через колонку "
            f"«Новая Розничная цена ЗЯ» в Excel — наценка новая пересчитается автоматически. "
            f"<br><br>Загрузите файл продаж, чтобы увидеть текущую валовую прибыль и эффект "
            f"выравнивания до медианы для позиций ниже рынка.",
        )
    else:
        delta = impact["delta_vp"]
        delta_cls = "ok" if delta >= 0 else "warn"
        _insight(
            delta_cls,
            f"Текущая средняя маржа <b>{_fmt_pct(base_m)}</b>. При выравнивании до медианы рынка (С2) "
            f"она составит <b>{_fmt_pct(c2_m)}</b>.<br><br>"
            f"По позициям <b>ниже медианы</b> с продажами ({impact['n_pos']} шт.): "
            f"текущая валовая прибыль <b>{_fmt_num(impact['current_vp'])}</b>. "
            f"Если поднять цену до медианы конкурентов — ВП станет "
            f"<b>{_fmt_num(impact['median_vp'])}</b> "
            f"(изменение <b>{_fmt_num(delta)}</b>). "
            f"Оценка только для позиций с ценой ЗЯ ниже медианной цены рынка.",
        )

# =============================== Вкладка: Данные ============================
with tab_data:
    if sales_df is not None:
        st.markdown("#### Продажи (сопоставление по штрихкоду)")
        st.caption(
            f"Загружено строк продаж: **{n_sales_rows}**, "
            f"совпало с прайсом: **{n_sales_matched}**. "
            "В Excel колонка «Кол-во продаж» на листе «Расчеты» заполняется через VLOOKUP."
        )
        sold = metrics[metrics["Есть_продажи"]].copy()
        if len(sold):
            sales_view = pd.DataFrame({
                "Наименование": sold["Наименование"].values,
                "Штрихкод": sold["Штрихкод"].values,
                "Кол-во продаж": sold["Кол-во_продаж"].apply(
                    lambda x: "—" if pd.isna(x) else f"{int(x):,}".replace(",", " ")
                ).values,
                "Выручка (файл)": sold["Выручка_продаж"].apply(_fmt_num).values,
                "ВП (файл)": sold["ВП_продаж"].apply(_fmt_num).values,
            })
            st.dataframe(sales_view, width="stretch", hide_index=True)
        else:
            st.warning(
                "Файл продаж загружен, но ни один штрихкод не совпал с прайсом. "
                "Проверьте колонку «Штрихкод»."
            )
        st.divider()

    st.markdown("#### Топ-10 позиций по текущей марже")
    tm = valid.nlargest(10, "Маржа")
    top_margin = pd.DataFrame({
        "Наименование": tm["Наименование"].values,
        "Маржа": tm["Маржа"].apply(_fmt_pct).values,
        "ВП/ед": tm["ВП_ед"].apply(_fmt_num).values,
        "Цена ЗЯ": tm["Розничная_цена_ЗЯ"].apply(_fmt_num).values,
    })
    st.dataframe(top_margin, width="stretch", hide_index=True)

    st.markdown("#### Топ-10 позиций выше рынка")
    ta = metrics[metrics["Сигнал"] == "Выше рынка"].nlargest(10, "Откл_рынок")
    if len(ta):
        top_above = pd.DataFrame({
            "Наименование": ta["Наименование"].values,
            "Откл. от рынка": ta["Откл_рынок"].apply(
                lambda x: "—" if pd.isna(x) else f"{x * 100:+.1f}%").values,
            "Цена ЗЯ": ta["Розничная_цена_ЗЯ"].apply(_fmt_num).values,
            "Ср. рынок": ta["Ср_рынок"].apply(_fmt_num).values,
        })
        st.dataframe(top_above, width="stretch", hide_index=True)
    else:
        st.caption("Позиций выше рынка не найдено.")

    st.divider()
    st.markdown("#### Итоговый Excel-файл")
    st.write(
        f"**{RELEASE_LABEL}.** Книга: Инструкция, Исходные данные, Продажи, Расчеты, "
        "Dashboard, Сценарный анализ. UI: объёмные KPI, ТОП отклонений с ценой конкурента, "
        "светофор выводов, сценарии С2/С4 и оценка ВП до медианы при продажах."
    )
    st.download_button(
        "⬇️ Скачать итоговый Excel (.xlsx)",
        data=excel_bytes,
        file_name="Зеленое_Яблоко_калькулятор.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
