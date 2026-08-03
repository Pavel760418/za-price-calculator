"""Построение листа 'Dashboard' - KPI-карточки, топы, светофоры."""
from __future__ import annotations

from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from za_price_calculator.config import PALETTE, SHEETS
from za_price_calculator.core.sheet_calculations import (
    COL_AVG_MKT,
    COL_DEV_AVG,
    COL_DEV_MED,
    COL_MARGIN_CUR,
    COL_MARGIN_SIG,
    COL_MARKUP_CUR,
    COL_RETAIL,
    COL_RISK,
    COL_SIGNAL,
    COL_VP_CUR,
)
from za_price_calculator.styling.styles import (
    HDR_ALIGN,
    NUM2,
    PCT,
    PCT_SIGNED,
    align,
    border_thin,
    fill,
    font,
)


def _L(col: int) -> str:
    return get_column_letter(col)


def _kpi(ws, row, col, label, formula, num_fmt="#,##0", bg=PALETTE.mid_blue):
    lc = ws.cell(row, col, label)
    lc.font = font(bold=True, size=9, color=PALETTE.white)
    lc.fill = fill(bg)
    lc.alignment = align("center")
    lc.border = border_thin()
    ws.row_dimensions[row].height = 22

    vc = ws.cell(row + 1, col, formula)
    vc.font = font(bold=True, size=18, color="1F3864")
    vc.fill = fill(PALETTE.light_blue)
    vc.alignment = align("center")
    vc.number_format = num_fmt
    vc.border = border_thin()
    ws.row_dimensions[row + 1].height = 34


def _section(ws, row, c1, c2, title, hexcolor=PALETTE.dark_blue):
    ws.merge_cells(f"{get_column_letter(c1)}{row}:{get_column_letter(c2)}{row}")
    c = ws.cell(row, c1, title)
    c.font = font(bold=True, size=11, color=PALETTE.white)
    c.fill = fill(hexcolor)
    c.alignment = align("left", indent=1)
    ws.row_dimensions[row].height = 22


def _top10_block(
    ws: Worksheet,
    *,
    section_row: int,
    header_row: int,
    c1: int,
    title: str,
    title_color: str,
    headers: list[str],
    header_color: str,
    cs: str,
    last: int,
    top_n: int,
    metric_col: str,
    retail_col: str,
    extreme: str,
) -> None:
    """
    Блок ТОП-10 по метрике отклонения.

    :param extreme: "LARGE" (выше рынка) или "SMALL" (ниже рынка).
    """
    c2, c3 = c1 + 1, c1 + 2
    _section(ws, section_row, c1, c3, title, title_color)
    for ci, h in enumerate(headers, start=c1):
        c = ws.cell(header_row, ci, h)
        c.fill = fill(header_color)
        c.font = font(bold=True, size=10, color=PALETTE.white)
        c.alignment = HDR_ALIGN
        c.border = border_thin()
    ws.row_dimensions[header_row].height = 18

    for rk in range(1, top_n + 1):
        row_i = header_row + rk
        ws.row_dimensions[row_i].height = 17
        filt = (
            f'IF({cs}!{metric_col}${3}:{metric_col}${last}<>"",'
            f'{cs}!{metric_col}${3}:{metric_col}${last})'
        )
        ext = f"{extreme}({filt},{rk})"

        c = ws.cell(
            row_i, c1,
            f'=IFERROR(INDEX({cs}!A$3:A${last},MATCH({ext},{cs}!{metric_col}$3:{metric_col}${last},0)),"")',
        )
        c.font = font(size=10)
        c.alignment = align("left", indent=1)
        c.border = border_thin()

        c = ws.cell(row_i, c2, f'=IFERROR({ext},"")')
        c.font = font(bold=True, size=10, color="C00000" if extreme == "LARGE" else "375623")
        c.number_format = PCT_SIGNED
        c.alignment = align("center")
        c.border = border_thin()

        c = ws.cell(
            row_i, c3,
            f'=IFERROR(INDEX({cs}!{retail_col}$3:{retail_col}${last},'
            f'MATCH({ext},{cs}!{metric_col}$3:{metric_col}${last},0)),"")',
        )
        c.font = font(size=10)
        c.number_format = NUM2
        c.alignment = align("right")
        c.border = border_thin()

    if top_n > 0:
        end_row = header_row + top_n
        letter = get_column_letter(c2)
        if extreme == "LARGE":
            ws.conditional_formatting.add(
                f"{letter}{header_row + 1}:{letter}{end_row}",
                ColorScaleRule(
                    start_type="min", start_color="FFEB84",
                    end_type="max", end_color="F8696B",
                ),
            )
        else:
            ws.conditional_formatting.add(
                f"{letter}{header_row + 1}:{letter}{end_row}",
                ColorScaleRule(
                    start_type="min", start_color="63BE7B",
                    end_type="max", end_color="FFEB84",
                ),
            )


def build_dashboard_sheet(ws: Worksheet, n_rows: int) -> None:
    """
    Строит лист 'Dashboard' с KPI-карточками и топ-10 списками.

    :param ws: целевой лист.
    :param n_rows: количество строк данных на листе 'Расчеты' (N).
    """
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = PALETTE.accent_green

    for col, w in (("A", 3), ("B", 28), ("C", 12), ("D", 12), ("E", 28), ("F", 12), ("G", 12), ("H", 3)):
        ws.column_dimensions[col].width = w

    cs = f"'{SHEETS.calculations}'"
    last = n_rows + 2 if n_rows else 3

    sig, mrg, mkup = _L(COL_SIGNAL), _L(COL_MARGIN_CUR), _L(COL_MARKUP_CUR)
    risk, vp, retail = _L(COL_RISK), _L(COL_VP_CUR), _L(COL_RETAIL)
    avg_mkt, dev_avg, dev_med = _L(COL_AVG_MKT), _L(COL_DEV_AVG), _L(COL_DEV_MED)
    msig = _L(COL_MARGIN_SIG)

    ws.merge_cells("B2:G2")
    t = ws["B2"]
    t.value = "ЗЕЛЕНОЕ ЯБЛОКО - Dashboard цен и маржи"
    t.font = font(bold=True, size=16, color=PALETTE.white)
    t.fill = fill(PALETTE.dark_blue)
    t.alignment = align("center")
    ws.row_dimensions[2].height = 36

    ws.merge_cells("B3:G3")
    s = ws["B3"]
    s.value = "Основные показатели по всему прайс-листу | Автообновление | Релиз 3.5"
    s.font = font(size=10, color="595959")
    s.alignment = align("center")
    ws.row_dimensions[3].height = 16

    if n_rows == 0:
        ws.merge_cells("B5:G6")
        c = ws["B5"]
        c.value = "Нет данных для расчёта KPI - загрузите прайс-лист с товарами."
        c.font = font(bold=True, size=12, color=PALETTE.accent_red)
        c.alignment = align("center", wrap=True)
        return

    kpis1 = [
        ("Всего позиций", f"=COUNTA({cs}!A3:A10000)", "#,##0", PALETTE.mid_blue),
        ("Выше рынка", f'=COUNTIF({cs}!{sig}3:{sig}10000,"Выше рынка")', "#,##0", PALETTE.brown),
        ("Ниже рынка", f'=COUNTIF({cs}!{sig}3:{sig}10000,"Ниже рынка")', "#,##0", PALETTE.dark_green),
        ("Средняя маржа", f'=IFERROR(AVERAGE(IF({cs}!{mrg}3:{mrg}{last}<>"",{cs}!{mrg}3:{mrg}{last})),"")', PCT, PALETTE.dark_blue),
        ("Средняя наценка", f'=IFERROR(AVERAGE(IF({cs}!{mkup}3:{mkup}{last}<>"",{cs}!{mkup}3:{mkup}{last})),"")', PCT, PALETTE.dark_blue),
        ("С критич.риском", f'=COUNTIF({cs}!{risk}3:{risk}10000,"КРИТИЧНО")', "#,##0", PALETTE.accent_red),
    ]
    for col_idx, (lbl, frm, nf, bg) in enumerate(kpis1, start=2):
        _kpi(ws, 5, col_idx, lbl, frm, nf, bg)

    ws.row_dimensions[8].height = 8

    kpis2 = [
        ("Ср.ВП/ед (расч.)", f'=IFERROR(AVERAGE(IF({cs}!{vp}3:{vp}{last}<>"",{cs}!{vp}3:{vp}{last})),"")', NUM2, PALETTE.mid_blue),
        ("Ср.цена ЗЯ", f'=IFERROR(AVERAGE(IF({cs}!{retail}3:{retail}{last}<>"",{cs}!{retail}3:{retail}{last})),"")', NUM2, PALETTE.mid_blue),
        ("Ср.цена рынка", f'=IFERROR(AVERAGE(IF({cs}!{avg_mkt}3:{avg_mkt}{last}<>"",{cs}!{avg_mkt}3:{avg_mkt}{last})),"")', NUM2, PALETTE.mid_blue),
        ("Ср.откл.от рынка", f'=IFERROR(AVERAGE(IF({cs}!{dev_avg}3:{dev_avg}{last}<>"",{cs}!{dev_avg}3:{dev_avg}{last})),"")', PCT_SIGNED, PALETTE.dark_blue),
        ("На уровне рынка", f'=COUNTIF({cs}!{sig}3:{sig}10000,"На уровне")', "#,##0", PALETTE.purple),
        ("Хор.маржа(>20%)", f'=COUNTIF({cs}!{msig}3:{msig}10000,"Хорошая(>20%)")', "#,##0", PALETTE.dark_green),
    ]
    for col_idx, (lbl, frm, nf, bg) in enumerate(kpis2, start=2):
        _kpi(ws, 9, col_idx, lbl, frm, nf, bg)

    ws.row_dimensions[12].height = 8

    # --- Топ-10 по марже (слева) ---
    _section(ws, 13, 2, 4, "Топ-10 позиций по текущей марже %")
    for ci, h in enumerate(["Наименование", "Маржа %", "ВП/ед"], start=2):
        c = ws.cell(14, ci, h)
        c.fill = fill(PALETTE.mid_blue)
        c.font = font(bold=True, size=10, color=PALETTE.white)
        c.alignment = HDR_ALIGN
        c.border = border_thin()
    ws.row_dimensions[14].height = 18

    top_n = min(10, n_rows)
    for rk in range(1, top_n + 1):
        row_i = 14 + rk
        ws.row_dimensions[row_i].height = 17
        c = ws.cell(row_i, 2,
            f'=IFERROR(INDEX({cs}!A$3:A${last},MATCH(LARGE(IF({cs}!{mrg}$3:{mrg}${last}<>"",'
            f'{cs}!{mrg}$3:{mrg}${last}),{rk}),{cs}!{mrg}$3:{mrg}${last},0)),"")')
        c.font = font(size=10)
        c.alignment = align("left", indent=1)
        c.border = border_thin()

        c = ws.cell(row_i, 3,
            f'=IFERROR(LARGE(IF({cs}!{mrg}$3:{mrg}${last}<>"",{cs}!{mrg}$3:{mrg}${last}),{rk}),"")')
        c.font = font(bold=True, size=10)
        c.number_format = PCT
        c.alignment = align("center")
        c.border = border_thin()

        c = ws.cell(row_i, 4,
            f'=IFERROR(INDEX({cs}!{vp}$3:{vp}${last},MATCH(LARGE(IF({cs}!{mrg}$3:{mrg}${last}<>"",'
            f'{cs}!{mrg}$3:{mrg}${last}),{rk}),{cs}!{mrg}$3:{mrg}${last},0)),"")')
        c.font = font(size=10)
        c.number_format = NUM2
        c.alignment = align("right")
        c.border = border_thin()

    if top_n > 0:
        ws.conditional_formatting.add(
            f"C15:C{14 + top_n}",
            ColorScaleRule(start_type="min", start_color="FFEB84", end_type="max", end_color="63BE7B"),
        )

    # Справка по критическому риску (справа от топа маржи)
    _section(ws, 13, 5, 7, "Что значит «критич. риск»", PALETTE.accent_red)
    ws.merge_cells("E14:G24")
    note = ws["E14"]
    note.value = (
        "«С критич.риском» — число позиций, у которых маржа упала бы ниже 5%, "
        "если снизить цену ЗЯ до минимального конкурента.\n\n"
        "КРИТИЧНО (<5%) — снижать цену опасно.\n"
        "УМЕРЕННЫЙ (5–15%) — требует мониторинга.\n"
        "НИЗКИЙ (>15%) — запас прочности достаточный.\n\n"
        "См. также столбец «Риск потери ВП» на листе Расчеты (светофор)."
    )
    note.font = font(size=10)
    note.alignment = align("left", wrap=True, indent=1)
    note.fill = fill(PALETTE.red_bg)
    note.border = border_thin()

    # --- Отклонения от средней ---
    ws.row_dimensions[25].height = 8
    _top10_block(
        ws, section_row=26, header_row=27, c1=2,
        title="Топ-10 выше рынка (откл. от средней)",
        title_color=PALETTE.brown, headers=["Наименование", "Откл.%", "Цена ЗЯ"],
        header_color=PALETTE.brown, cs=cs, last=last, top_n=top_n,
        metric_col=dev_avg, retail_col=retail, extreme="LARGE",
    )
    _top10_block(
        ws, section_row=26, header_row=27, c1=5,
        title="Топ-10 ниже рынка (откл. от средней)",
        title_color=PALETTE.dark_green, headers=["Наименование", "Откл.%", "Цена ЗЯ"],
        header_color=PALETTE.dark_green, cs=cs, last=last, top_n=top_n,
        metric_col=dev_avg, retail_col=retail, extreme="SMALL",
    )

    # --- Отклонения от медианы ---
    ws.row_dimensions[38].height = 8
    _top10_block(
        ws, section_row=39, header_row=40, c1=2,
        title="Топ-10 выше рынка (откл. от медианы)",
        title_color=PALETTE.brown, headers=["Наименование", "Откл.%", "Цена ЗЯ"],
        header_color=PALETTE.brown, cs=cs, last=last, top_n=top_n,
        metric_col=dev_med, retail_col=retail, extreme="LARGE",
    )
    _top10_block(
        ws, section_row=39, header_row=40, c1=5,
        title="Топ-10 ниже рынка (откл. от медианы)",
        title_color=PALETTE.dark_green, headers=["Наименование", "Откл.%", "Цена ЗЯ"],
        header_color=PALETTE.dark_green, cs=cs, last=last, top_n=top_n,
        metric_col=dev_med, retail_col=retail, extreme="SMALL",
    )
